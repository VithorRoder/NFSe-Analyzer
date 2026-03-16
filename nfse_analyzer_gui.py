from pathlib import Path
from playwright.sync_api import sync_playwright
from decimal import Decimal, InvalidOperation
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from datetime import datetime
import hashlib
import re
import sys
import uuid
import traceback
import tkinter as tk
from tkinter import ttk, messagebox

SECRET = "NFSE_ANALYZER_2026_RODER"
URL_PORTAL = "https://www.nfse.gov.br/EmissorNacional"


def get_base_dir() -> Path:
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent


BASE_DIR = get_base_dir()
ARQUIVO_XLSX = BASE_DIR / "notas_emitidas_resumo.xlsx"


def formatar_machine_id(node: int) -> str:
    return f"{node:012X}"


def get_machine_id() -> str:
    return formatar_machine_id(uuid.getnode())


def gerar_hash_licenca(machine_id: str) -> str:
    texto = machine_id.strip().upper() + SECRET
    return hashlib.sha256(texto.encode("utf-8")).hexdigest()


def verificar_licenca() -> Path:
    caminhos_possiveis = [
        BASE_DIR / "licenca.key",
    ]
    machine_id = get_machine_id()

    licenca_path = next((p for p in caminhos_possiveis if p.exists()), None)

    if licenca_path is None:
        raise FileNotFoundError(
            "Licença não encontrada.\n\n"
            f"Coloque o arquivo licenca.key na mesma pasta do programa:\n{BASE_DIR}\n\n"
            f"ID desta máquina para ativação:\n{machine_id}"
        )

    licenca_informada = licenca_path.read_text(encoding="utf-8").strip()
    licenca_valida = gerar_hash_licenca(machine_id)

    if licenca_informada != licenca_valida:
        raise PermissionError(
            "Licença inválida para este computador.\n\n"
            f"Arquivo lido: {licenca_path}\n"
            f"ID desta máquina: {machine_id}"
        )

    return licenca_path


def limpar_espacos(s: str) -> str:
    return re.sub(r"\s+", " ", s or "").strip()


def formatar_brl_decimal(valor: Decimal) -> str:
    s = f"{valor:.2f}"
    inteiro, decimal = s.split(".")
    grupos = []
    while len(inteiro) > 3:
        grupos.insert(0, inteiro[-3:])
        inteiro = inteiro[:-3]
    grupos.insert(0, inteiro)
    return f"R$ {'.'.join(grupos)},{decimal}"


def parse_valor_brl(texto: str) -> Decimal:
    if not texto:
        return Decimal("0")

    texto = texto.strip()
    encontrados = re.findall(r"\d{1,3}(?:\.\d{3})*,\d{2}", texto)
    if not encontrados:
        return Decimal("0")

    valor = encontrados[-1].replace(".", "").replace(",", ".")
    try:
        return Decimal(valor)
    except InvalidOperation:
        return Decimal("0")


def obter_pagina_atual(page) -> int:
    url = page.url
    m = re.search(r"[?&]pg=(\d+)", url)
    if m:
        return int(m.group(1))
    return 1


def obter_linhas_da_tabela(page):
    seletores = ["table tbody tr", "tbody tr", "table tr"]
    for seletor in seletores:
        try:
            linhas = page.locator(seletor)
            count = linhas.count()
            if count > 0:
                return linhas, count, seletor
        except Exception:
            pass
    return None, 0, None


def linha_tem_icone_cancelada(row) -> bool:
    seletores = [
        "[title*='cancel' i]",
        "[aria-label*='cancel' i]",
        "[data-bs-original-title*='cancel' i]",
        "[data-original-title*='cancel' i]",
        "i[title*='cancel' i]",
        "svg[title*='cancel' i]",
        "img[title*='cancel' i]",
    ]

    for seletor in seletores:
        try:
            itens = row.locator(seletor)
            if itens.count() > 0:
                return True
        except Exception:
            pass

    try:
        icones = row.locator("i, svg, img, span")
        qtd_icones = icones.count()
        for i in range(qtd_icones):
            el = icones.nth(i)
            attrs = []
            for attr in ["title", "aria-label", "data-bs-original-title", "data-original-title", "class"]:
                try:
                    v = el.get_attribute(attr)
                except Exception:
                    v = None
                if v:
                    attrs.append(v.lower())

            if "cancel" in " ".join(attrs):
                return True
    except Exception:
        pass

    return False


def identificar_cancelada(texto: str, situacao: str) -> bool:
    base = f"{texto} {situacao}".lower()
    termos_cancelada = [
        "cancelada",
        "cancelado",
        "nfse cancelada",
        "nfs-e cancelada",
        "situação: cancelada",
        "situacao: cancelada",
        "status: cancelada",
    ]
    return any(termo in base for termo in termos_cancelada)


def extrair_dados_linha_por_colunas(row):
    data = ""
    cliente = ""
    descricao = ""
    situacao = ""
    valor = Decimal("0")

    try:
        colunas = row.locator("td")
        qtd_colunas = colunas.count()
    except Exception:
        qtd_colunas = 0

    valores_colunas = []
    for idx in range(qtd_colunas):
        try:
            txt = limpar_espacos(colunas.nth(idx).text_content() or "")
        except Exception:
            txt = ""
        valores_colunas.append(txt)

    texto_completo = " ".join(v for v in valores_colunas if v)

    if qtd_colunas >= 1:
        data = valores_colunas[0]
    if qtd_colunas >= 2:
        cliente = valores_colunas[1]
    if qtd_colunas >= 3:
        descricao = valores_colunas[2]
    if qtd_colunas >= 4:
        situacao = valores_colunas[3]

    for txt in reversed(valores_colunas):
        valor_extraido = parse_valor_brl(txt)
        if valor_extraido > 0:
            valor = valor_extraido
            break

    return data, cliente, descricao, situacao, valor, texto_completo


def coletar_notas_da_pagina(page, pagina, vistos_textos, log):
    registros = []

    linhas, qtd, seletor = obter_linhas_da_tabela(page)
    if not linhas or qtd == 0:
        log(f"[AVISO] Nenhuma linha encontrada na página {pagina}.")
        return registros

    log(f"[INFO] Página {pagina}: {qtd} linhas encontradas com seletor {seletor}")

    for i in range(qtd):
        row = linhas.nth(i)
        data, cliente, descricao, situacao, valor, texto = extrair_dados_linha_por_colunas(row)

        if not texto:
            continue
        if not re.search(r"\d{2}/\d{2}/\d{4}", texto):
            continue

        chave = f"{pagina}-{i}"
        if chave in vistos_textos:
            continue
        vistos_textos.add(chave)

        cancelada = identificar_cancelada(texto, situacao)
        if not cancelada:
            cancelada = linha_tem_icone_cancelada(row)

        dados = {
            "pagina": pagina,
            "data": data,
            "cliente": cliente,
            "descricao": descricao,
            "situacao": situacao,
            "valor": valor,
            "cancelada": cancelada,
            "texto_original": texto,
        }

        status_txt = "CANCELADA" if cancelada else "OK"
        log(
            f"[NOTA] pág={pagina} status={status_txt} "
            f"data={dados['data']} valor={formatar_brl_decimal(valor)} cliente={dados['cliente']}"
        )

        registros.append(dados)

    return registros


def ir_para_proxima_pagina(page, pagina_atual):
    proxima = pagina_atual + 1

    try:
        link = page.locator(f"a[href*='pg={proxima}']").first
        if link.count() > 0:
            link.click(timeout=5000)
            page.wait_for_load_state("domcontentloaded", timeout=15000)
            page.wait_for_timeout(500)
            return True
    except Exception:
        pass

    return False


def aplicar_larguras(ws) -> None:
    larguras = {}
    for row in ws.iter_rows():
        for cell in row:
            valor = "" if cell.value is None else str(cell.value)
            larguras[cell.column] = max(larguras.get(cell.column, 0), len(valor))

    for idx, largura in larguras.items():
        ws.column_dimensions[get_column_letter(idx)].width = min(max(largura + 2, 12), 60)


def estilizar_cabecalho(ws) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center")


def criar_xlsx(todos_registros) -> None:
    wb = Workbook()

    todos_registros = sorted(
        todos_registros,
        key=lambda r: datetime.strptime(r["data"], "%d/%m/%Y")
    )

    ws_resumo = wb.active
    ws_resumo.title = "Resumo"

    registros_validos = [r for r in todos_registros if not r["cancelada"]]
    registros_cancelados = [r for r in todos_registros if r["cancelada"]]

    total_validas = sum((r["valor"] for r in registros_validos), Decimal("0"))
    total_canceladas = sum((r["valor"] for r in registros_cancelados), Decimal("0"))

    ws_resumo.append(["Indicador", "Valor"])
    ws_resumo.append(["Quantidade total lida", len(todos_registros)])
    ws_resumo.append(["Notas válidas", len(registros_validos)])
    ws_resumo.append(["Notas canceladas", len(registros_cancelados)])
    ws_resumo.append(["Total notas válidas", float(total_validas)])
    ws_resumo.append(["Total notas canceladas", float(total_canceladas)])

    estilizar_cabecalho(ws_resumo)

    ws_resumo["B5"].number_format = 'R$ * #,##0.00'
    ws_resumo["B6"].number_format = 'R$ * #,##0.00'

    ws_notas = wb.create_sheet("Notas")
    ws_notas.append(["Página", "Data", "Cliente", "Competência", "Município Emissor", "Cancelada", "Valor"])
    estilizar_cabecalho(ws_notas)

    for r in todos_registros:
        ws_notas.append([
            r["pagina"],
            r["data"],
            r["cliente"],
            r["descricao"],
            r["situacao"],
            "SIM" if r["cancelada"] else "NÃO",
            float(r["valor"]),
        ])

    for linha in range(2, ws_notas.max_row + 1):
        ws_notas[f"G{linha}"].number_format = 'R$ #,##0.00'
        if ws_notas[f"F{linha}"].value == "SIM":
            for cell in ws_notas[linha]:
                cell.fill = PatternFill("solid", fgColor="FCE4D6")

    aplicar_larguras(ws_resumo)
    aplicar_larguras(ws_notas)

    ws_resumo.column_dimensions["B"].width = 19
    ws_notas.column_dimensions["C"].width = 90
    ws_notas.column_dimensions["G"].width = 19

    wb.save(ARQUIVO_XLSX)


def abrir_browser_do_sistema(playwright):
    import tkinter as tk

    root = tk.Tk()
    root.withdraw()
    largura = root.winfo_screenwidth()
    altura = root.winfo_screenheight()
    root.destroy()

    try:
        print("Abrindo Google Chrome...")
        browser = playwright.chromium.launch(
            channel="chrome",
            headless=False,
            args=[
                f"--window-size={largura},{altura}",
                "--window-position=0,0",
            ],
        )
        navegador = "Chrome"
    except Exception:
        try:
            print("Chrome não encontrado. Abrindo Microsoft Edge...")
            browser = playwright.chromium.launch(
                channel="msedge",
                headless=False,
                args=[
                    f"--window-size={largura},{altura}",
                    "--window-position=0,0",
                ],
            )
            navegador = "Edge"
        except Exception:
            raise RuntimeError(
                "Nenhum navegador compatível encontrado.\n\n"
                "Instale Google Chrome ou Microsoft Edge."
            )

    context = browser.new_context(
        viewport={"width": largura, "height": altura}
    )

    return browser, context, navegador


class NFSeAnalyzerApp(tk.Tk):
    def __init__(self) -> None:
        super().__init__()
        self.title("NFSe Analyzer")
        self.geometry("900x620")
        self.minsize(820, 560)

        self.var_status = tk.StringVar(value="Pronto para iniciar.")
        self.var_arquivo = tk.StringVar(value=str(ARQUIVO_XLSX))
        self.var_base_dir = tk.StringVar(value=str(BASE_DIR))
        self.var_machine_id = tk.StringVar(value=get_machine_id())
        self.var_licenca = tk.StringVar(value="Ainda não verificada")

        self.text_log = None
        self.botao_iniciar = None
        self._montar_interface()
        self._verificar_licenca_inicial()

    def _montar_interface(self) -> None:
        main = ttk.Frame(self, padding=14)
        main.pack(fill="both", expand=True)

        ttk.Label(main, text="NFSe Analyzer", font=("Segoe UI", 16, "bold")).pack(anchor="w")
        ttk.Label(
            main,
            text="Leitura de notas emitidas do portal nacional da NFS-e usando o navegador do sistema.",
            wraplength=860
        ).pack(anchor="w", pady=(4, 12))

        info = ttk.LabelFrame(main, text="Informações", padding=10)
        info.pack(fill="x")

        self._linha_info(info, "Pasta do programa:", self.var_base_dir)
        self._linha_info(info, "Arquivo Excel:", self.var_arquivo)
        self._linha_info(info, "ID da máquina:", self.var_machine_id)
        self._linha_info(info, "Licença:", self.var_licenca)

        botoes = ttk.Frame(main)
        botoes.pack(fill="x", pady=12)

        self.botao_iniciar = ttk.Button(botoes, text="Abrir portal e coletar notas", command=self.iniciar)
        self.botao_iniciar.pack(side="left")
        ttk.Button(botoes, text="Verificar licença novamente", command=self.verificar_licenca_manual).pack(side="left", padx=(8, 0))
        ttk.Button(botoes, text="Fechar", command=self.destroy).pack(side="right")

        ttk.Label(main, text="Log de execução:").pack(anchor="w", pady=(6, 4))

        self.text_log = tk.Text(main, height=20, wrap="word")
        self.text_log.pack(fill="both", expand=True)
        self.text_log.configure(state="disabled")

        status = ttk.Label(main, textvariable=self.var_status)
        status.pack(anchor="w", pady=(10, 0))

    def _linha_info(self, parent, titulo: str, var: tk.StringVar) -> None:
        row = ttk.Frame(parent)
        row.pack(fill="x", pady=2)
        ttk.Label(row, text=titulo, width=18).pack(side="left")
        ttk.Entry(row, textvariable=var, state="readonly").pack(side="left", fill="x", expand=True)

    def log(self, msg: str) -> None:
        self.text_log.configure(state="normal")
        self.text_log.insert("end", msg + "\n")
        self.text_log.see("end")
        self.text_log.configure(state="disabled")
        self.update_idletasks()

    def set_status(self, msg: str) -> None:
        self.var_status.set(msg)
        self.update_idletasks()

    def _verificar_licenca_inicial(self) -> None:
        try:
            licenca = verificar_licenca()
            self.var_licenca.set(f"OK - {licenca.name}")
            self.log(f"Licença carregada de: {licenca}")
        except Exception as exc:
            self.var_licenca.set("Ausente ou inválida")
            self.log(str(exc))
            messagebox.showwarning("Licença", str(exc))

    def verificar_licenca_manual(self) -> None:
        try:
            licenca = verificar_licenca()
            self.var_licenca.set(f"OK - {licenca.name}")
            messagebox.showinfo("Licença", f"Licença válida.\n\nArquivo: {licenca}")
            self.log(f"Licença validada novamente: {licenca}")
        except Exception as exc:
            self.var_licenca.set("Ausente ou inválida")
            messagebox.showerror("Licença", str(exc))
            self.log(str(exc))

    def iniciar(self) -> None:
        try:
            licenca = verificar_licenca()
            self.var_licenca.set(f"OK - {licenca.name}")
        except Exception as exc:
            self.var_licenca.set("Ausente ou inválida")
            messagebox.showerror("Licença", str(exc))
            self.log(str(exc))
            return

        self.botao_iniciar.configure(state="disabled")
        self.set_status("Executando...")

        try:
            self.executar_fluxo()
        except Exception as exc:
            self.log(traceback.format_exc())
            messagebox.showerror("Erro na execução", str(exc))
            self.set_status("Erro na execução.")
        finally:
            self.botao_iniciar.configure(state="normal")

    def executar_fluxo(self) -> None:
        self.log("Iniciando abertura do portal...")

        with sync_playwright() as p:
            browser, context, nome_browser = abrir_browser_do_sistema(p)
            page = context.new_page()

            self.log(f"Navegador aberto: {nome_browser}")
            page.goto(URL_PORTAL, wait_until="domcontentloaded")
            self.log("Portal da NFS-e aberto.")

            messagebox.showinfo(
                "Passo 1",
                "Faça login no portal da NFS-e no navegador que foi aberto.\n\n"
                "Depois clique em OK para continuar."
            )

            ok = messagebox.askokcancel(
                "Passo 2",
                "Agora abra a tela de NOTAS EMITIDAS e aplique o filtro desejado.\n\n"
                "Clique em OK quando a listagem estiver pronta para leitura."
            )
            if not ok:
                browser.close()
                self.set_status("Operação cancelada pelo usuário.")
                self.log("Operação cancelada pelo usuário antes da coleta.")
                return

            todos_registros = []
            vistos_textos = set()
            paginas_processadas = set()

            while True:
                pagina_atual = obter_pagina_atual(page)
                if pagina_atual in paginas_processadas:
                    break

                paginas_processadas.add(pagina_atual)
                registros = coletar_notas_da_pagina(page, pagina_atual, vistos_textos, self.log)
                todos_registros.extend(registros)

                if not ir_para_proxima_pagina(page, pagina_atual):
                    break

            if not todos_registros:
                browser.close()
                self.set_status("Nenhuma nota encontrada.")
                messagebox.showwarning("Resultado", "Nenhuma nota foi encontrada na tela atual.")
                self.log("Nenhuma nota encontrada para gerar o Excel.")
                return

            criar_xlsx(todos_registros)

            registros_validos = [r for r in todos_registros if not r["cancelada"]]
            canceladas = [r for r in todos_registros if r["cancelada"]]
            total = sum((r["valor"] for r in registros_validos), Decimal("0"))

            resumo = (
                f"Quantidade total lida: {len(todos_registros)}\n"
                f"Notas válidas: {len(registros_validos)}\n"
                f"Notas canceladas: {len(canceladas)}\n"
                f"Valor total (sem canceladas): {formatar_brl_decimal(total)}\n\n"
                f"Excel gerado em:\n{ARQUIVO_XLSX}"
            )

            self.log("=" * 80)
            self.log("RESUMO FINAL")
            self.log(resumo)
            self.log("=" * 80)

            browser.close()
            self.set_status("Concluído com sucesso.")
            messagebox.showinfo("Concluído", resumo)


if __name__ == "__main__":
    app = NFSeAnalyzerApp()
    app.mainloop()
