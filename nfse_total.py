
from pathlib import Path
from playwright.sync_api import sync_playwright
from decimal import Decimal, InvalidOperation
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from datetime import datetime
import re
import uuid
import hashlib
import time
import sys

SECRET = "NFSE_ANALYZER_2026_RODER"


def get_base_dir():
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent


def get_machine_id():
    return hex(uuid.getnode())


def gerar_hash_licenca(machine_id: str) -> str:
    texto = machine_id + SECRET
    return hashlib.sha256(texto.encode("utf-8")).hexdigest()


def verificar_licenca():
    base_dir = get_base_dir()
    licenca_path = base_dir / "licenca.key"

    machine_id = get_machine_id()

    if not licenca_path.exists():
        print("\n" + "=" * 80)
        print("LICENÇA NÃO ENCONTRADA")
        print("=" * 80)
        print("Envie este ID da máquina para ativação:\n")
        print(machine_id)
        input("\nPressione ENTER para sair...")
        sys.exit()

    licenca_informada = licenca_path.read_text(encoding="utf-8").strip()
    licenca_valida = gerar_hash_licenca(machine_id)

    if licenca_informada != licenca_valida:
        print("\n" + "=" * 80)
        print("LICENÇA INVÁLIDA")
        print("=" * 80)
        print("Esta licença não pertence a este computador.\n")
        print("ID da máquina:", machine_id)
        input("\nPressione ENTER para sair...")
        sys.exit()

if getattr(sys, "frozen", False):
    BASE_DIR = Path(sys.executable).resolve().parent
else:
    BASE_DIR = Path(__file__).resolve().parent

ARQUIVO_XLSX = BASE_DIR / "notas_emitidas_resumo.xlsx"

URL_PORTAL = "https://www.nfse.gov.br/EmissorNacional"


def esperar_usuario(msg: str):
    print("\n" + "=" * 80)
    print(msg)
    print("=" * 80)
    input("Pressione ENTER para continuar... ")


def limpar_espacos(s: str) -> str:
    return re.sub(r"\s+", " ", s or "").strip()


def formatar_brl_decimal(valor: Decimal) -> str:
    s = f"{valor:.2f}"
    inteiro, decimal = s.split(".")
    grupos = []
    while len(inteiro) > 3:
        grupos.insert(0, inteiro[-3:])
        inteiro = inteiro[:-3]
    grupos.insert(0, inteiro)
    return f"R$ {'.'.join(grupos)},{decimal}"


def parse_valor_brl(texto: str) -> Decimal:
    if not texto:
        return Decimal("0")

    texto = texto.strip()
    encontrados = re.findall(r"\d{1,3}(?:\.\d{3})*,\d{2}", texto)
    if not encontrados:
        return Decimal("0")

    valor = encontrados[-1].replace(".", "").replace(",", ".")
    try:
        return Decimal(valor)
    except InvalidOperation:
        return Decimal("0")


def obter_pagina_atual(page):
    url = page.url
    m = re.search(r"[?&]pg=(\d+)", url)
    if m:
        return int(m.group(1))
    return 1


def obter_linhas_da_tabela(page):
    seletores = ["table tbody tr", "tbody tr", "table tr"]
    for seletor in seletores:
        try:
            linhas = page.locator(seletor)
            count = linhas.count()
            if count > 0:
                return linhas, count, seletor
        except Exception:
            pass
    return None, 0, None


def linha_tem_icone_cancelada(row) -> bool:
    seletores = [
        "[title*='cancel' i]",
        "[aria-label*='cancel' i]",
        "[data-bs-original-title*='cancel' i]",
        "[data-original-title*='cancel' i]",
        "i[title*='cancel' i]",
        "svg[title*='cancel' i]",
        "img[title*='cancel' i]",
    ]

    for seletor in seletores:
        try:
            itens = row.locator(seletor)
            if itens.count() > 0:
                return True
        except Exception:
            pass

    try:
        icones = row.locator("i, svg, img, span")
        for i in range(icones.count()):
            el = icones.nth(i)

            attrs = []
            for attr in ["title", "aria-label", "data-bs-original-title", "data-original-title", "class"]:
                try:
                    v = el.get_attribute(attr)
                except Exception:
                    v = None
                if v:
                    attrs.append(v.lower())

            texto_attrs = " ".join(attrs)
            if "cancel" in texto_attrs:
                return True
    except Exception:
        pass

    return False


def identificar_cancelada(texto: str, situacao: str, row=None) -> bool:
    base = f"{texto} {situacao}".lower()

    termos_cancelada = [
        "cancelada",
        "cancelado",
        "nfse cancelada",
        "nfs-e cancelada",
        "situação: cancelada",
        "situacao: cancelada",
        "status: cancelada",
    ]

    if any(termo in base for termo in termos_cancelada):
        return True

    if row is not None and linha_tem_icone_cancelada(row):
        return True

    return False


def extrair_dados_linha_por_colunas(row):
    data = ""
    cliente = ""
    descricao = ""
    situacao = ""
    valor = Decimal("0")

    try:
        colunas = row.locator("td")
        qtd_colunas = colunas.count()
    except Exception:
        qtd_colunas = 0

    valores_colunas = []
    for idx in range(qtd_colunas):
        try:
            txt = limpar_espacos(colunas.nth(idx).inner_text(timeout=1000))
        except Exception:
            txt = ""
        valores_colunas.append(txt)

    if qtd_colunas >= 1:
        data = valores_colunas[0]

    if qtd_colunas >= 2:
        cliente = valores_colunas[1]

    if qtd_colunas >= 3:
        descricao = valores_colunas[2]

    if qtd_colunas >= 4:
        situacao = valores_colunas[3]

    for txt in reversed(valores_colunas):
        valor_extraido = parse_valor_brl(txt)
        if valor_extraido > 0:
            valor = valor_extraido
            break

    return data, cliente, descricao, situacao, valor


def coletar_notas_da_pagina(page, pagina, vistos_textos):
    registros = []

    linhas, qtd, seletor = obter_linhas_da_tabela(page)
    if not linhas or qtd == 0:
        print(f"[AVISO] Nenhuma linha encontrada na página {pagina}.")
        return registros

    print(f"[INFO] Página {pagina}: {qtd} linhas encontradas com seletor {seletor}")

    for i in range(qtd):
        row = linhas.nth(i)

        try:
            texto = limpar_espacos(row.inner_text(timeout=1500))
        except Exception:
            texto = ""

        if not texto:
            continue

        if not re.search(r"\d{2}/\d{2}/\d{4}", texto):
            continue

        chave = f"{pagina}-{i}"
        if chave in vistos_textos:
            continue
        vistos_textos.add(chave)

        data, cliente, descricao, situacao, valor = extrair_dados_linha_por_colunas(row)

        cancelada = identificar_cancelada(texto, situacao, row)

        dados = {
            "pagina": pagina,
            "data": data,
            "cliente": cliente,
            "descricao": descricao,
            "situacao": situacao,
            "valor": valor,
            "cancelada": cancelada,
            "texto_original": texto,
        }

        status_txt = "CANCELADA" if cancelada else "OK"

        print(
            f"[NOTA] pág={pagina} status={status_txt} "
            f"data={dados['data']} "
            f"valor={formatar_brl_decimal(valor)} "
            f"cliente={dados['cliente']}"
        )

        registros.append(dados)

    return registros


def ir_para_proxima_pagina(page, pagina_atual):
    proxima = pagina_atual + 1

    try:
        link = page.locator(f"a[href*='pg={proxima}']").first
        if link.count() > 0:
            href = link.get_attribute("href")
            print(f"[PAGINAÇÃO] Indo para página {proxima} via href={href}")
            link.click(timeout=5000)
            page.wait_for_load_state("domcontentloaded", timeout=15000)
            time.sleep(2)
            return True
    except Exception:
        pass

    return False


def aplicar_larguras(ws):
    larguras = {}
    for row in ws.iter_rows():
        for cell in row:
            valor = "" if cell.value is None else str(cell.value)
            larguras[cell.column] = max(larguras.get(cell.column, 0), len(valor))

    for idx, largura in larguras.items():
        ws.column_dimensions[get_column_letter(idx)].width = min(max(largura + 2, 12), 60)


def estilizar_cabecalho(ws):
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center")


def criar_xlsx(todos_registros):
    wb = Workbook()

    todos_registros = sorted(
        todos_registros,
        key=lambda r: datetime.strptime(r["data"], "%d/%m/%Y")
    )

    ws_resumo = wb.active
    ws_resumo.title = "Resumo"

    registros_validos = [r for r in todos_registros if not r["cancelada"]]
    registros_cancelados = [r for r in todos_registros if r["cancelada"]]

    total_validas = sum((r["valor"] for r in registros_validos), Decimal("0"))
    total_canceladas = sum((r["valor"] for r in registros_cancelados), Decimal("0"))

    ws_resumo.append(["Indicador", "Valor"])
    ws_resumo.append(["Quantidade total lida", len(todos_registros)])
    ws_resumo.append(["Notas válidas", len(registros_validos)])
    ws_resumo.append(["Notas canceladas", len(registros_cancelados)])
    ws_resumo.append(["Total notas válidas", float(total_validas)])
    ws_resumo.append(["Total notas canceladas", float(total_canceladas)])

    estilizar_cabecalho(ws_resumo)

    ws_resumo["B5"].number_format = 'R$ * #,##0.00'
    ws_resumo["B6"].number_format = 'R$ * #,##0.00'

    ws_notas = wb.create_sheet("Notas")
    ws_notas.append(["Página", "Data", "Cliente", "Competência", "Município Emissor", "Cancelada", "Valor"])
    estilizar_cabecalho(ws_notas)

    for r in todos_registros:
        ws_notas.append([
            r["pagina"],
            r["data"],
            r["cliente"],
            r["descricao"],
            r["situacao"],
            "SIM" if r["cancelada"] else "NÃO",
            float(r["valor"])
        ])

    for linha in range(2, ws_notas.max_row + 1):
        ws_notas[f"G{linha}"].number_format = 'R$ #,##0.00'

        if ws_notas[f"F{linha}"].value == "SIM":
            for cell in ws_notas[linha]:
                cell.fill = PatternFill("solid", fgColor="FCE4D6")

    aplicar_larguras(ws_resumo)
    aplicar_larguras(ws_notas)

    ws_resumo.column_dimensions["B"].width = 19   
    ws_notas.column_dimensions["C"].width = 90   
    ws_notas.column_dimensions["G"].width = 19 

    wb.save(ARQUIVO_XLSX)


def main():
    verificar_licenca()
    with sync_playwright() as p:
        browser = p.chromium.launch(
        headless=False,
        channel="chrome"
        )
        context = browser.new_context()
        page = context.new_page()

        print("Abrindo portal da NFS-e...")
        page.goto(URL_PORTAL, wait_until="domcontentloaded")

        esperar_usuario("Faça login no portal da NFS-e.")

        esperar_usuario("Abra a tela de NOTAS EMITIDAS e filtre o período desejado.")

        todos_registros = []
        vistos_textos = set()
        paginas_processadas = set()

        while True:
            pagina_atual = obter_pagina_atual(page)

            if pagina_atual in paginas_processadas:
                break

            paginas_processadas.add(pagina_atual)

            registros = coletar_notas_da_pagina(page, pagina_atual, vistos_textos)
            todos_registros.extend(registros)

            avancou = ir_para_proxima_pagina(page, pagina_atual)
            if not avancou:
                break

        criar_xlsx(todos_registros)

        registros_validos = [r for r in todos_registros if not r["cancelada"]]
        canceladas = [r for r in todos_registros if r["cancelada"]]
        total = sum((r["valor"] for r in registros_validos), Decimal("0"))

        print("\n" + "=" * 80)
        print("RESUMO FINAL")
        print("=" * 80)
        print(f"Quantidade total lida: {len(todos_registros)}")
        print(f"Notas válidas: {len(registros_validos)}")
        print(f"Notas canceladas: {len(canceladas)}")
        print(f"Valor total (sem canceladas): {formatar_brl_decimal(total)}")
        print(f"Excel gerado em: {ARQUIVO_XLSX}")

        browser.close()


if __name__ == "__main__":
    main()
